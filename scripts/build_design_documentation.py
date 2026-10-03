"""Build docs/design documentation.pdf from the Harborline design write-up."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_pdfs import build_pdf, wrap  # noqa: E402

OUT = ROOT / "docs" / "design documentation.pdf"
DOCX = ROOT / "docs" / "design documentation.docx"

BODY = """
HARBORLINE
DESIGN DOCUMENTATION

This note records how Harborline answers employee questions: a hand-written
orchestrator, one MCP server, and OpenAI embeddings with an in-process cosine
index. The two People Desk demos follow fixed tool sequences.

DESIGN CHOICES

Orchestration. harborline/agent.py classifies intent with regular expressions
and then runs a fixed workflow. The model does not choose tools. That keeps HR
paths repeatable: the same question always calls the same tools, and the trace
lists each name, argument, and result. When HARBORLINE_ANSWER_MODE=llm and
OPENAI_API_KEY is set, the chat model only rewrites the draft after the tools
return. Refusals, missing records, and thin evidence stay on that draft.

MCP server. One FastMCP server, named harborline, exposes eight tools. Policy
tools read the RAG index. Profile, PTO, and benefits tools read mock HarborHub
JSON under data/. Ticket and email tools are session-only mocks and do not
write data/tickets.json. The agent discovers tools with tools/list and calls
only names from that list through tools/call. FastMCP builds each JSON Schema
from the Python signature, so the schema stays tied to the function.

Tool schemas.
- search_policy_documents: query (required), employee_id, kind
- get_policy_section: policy_id (required), section
- lookup_employee_profile: employee_id
- check_pto_balance: employee_id
- lookup_benefits_status: employee_id
- create_mock_hr_ticket: topic, summary, employee_id, confirm
- draft_hr_email: employee_id, subject, body, confirm
- check_policy_compliance: scenario, employee_id, policy_id

Transport. The CLI and Cursor use stdio: the client spawns
python -m harborline.mcp_server and speaks JSON-RPC on stdin and stdout, with
no open port. The People Desk page, POST /chat, and POST /demos/{id} use
mcp-inproc, the same FastMCP object inside the web process, so a Render
container does not spawn a second Python process per request. Streamable HTTP
on 127.0.0.1:8765 is available when an external client needs an HTTP MCP
endpoint.

Embeddings and chunking. Embeddings are OpenAI text-embedding-3-small
(1536 dimensions) from the embeddings API. Ingest splits each handbook by
heading. A section longer than 900 characters is cut into windows of 900 with
120 characters of overlap, so a policy fact stays with its heading and
neighboring sentences still share context.

Retrieval k and index. Search embeds the rewritten query, takes a candidate
pool of HARBORLINE_FETCH_K=20, then reranks to HARBORLINE_TOP_K=5. The blend
is 0.65 cosine and 0.35 lexical overlap, and the reranker prefers distinct
source files. The score floor is HARBORLINE_MIN_SCORE=0.22. Vectors are
L2-normalized and cached under .cache/. Query time ranks that matrix in the
process. tfidf remains the offline path used by pytest and CI.

Deployment. A Docker image runs FastAPI with uvicorn on port 8000. The image
copies corpus/, data/, and eval/ and does not download local embedding
weights. On Render, set OPENAI_API_KEY and HARBORLINE_ANSWER_MODE=llm. The
Dockerfile pins retrieve, so the service variable overrides it. The first
search after a restart embeds the corpus and holds the matrix in that
process. The cache does not survive a new deploy.

Safety guardrails. Out-of-corpus questions (stock tips, other employers,
medical advice, and similar patterns) are refused. Empty retrieval or a top
score under 0.22 returns a weak-match refusal. Missing employee ids ask for
clarification. Missing HarborHub rows escalate. Mock ticket and email writes
stay pending until confirm=true, and even then they exist only in that
process. The model rewrite leaves refusals in place for an unavailable tool
bus, a missing employee, or incomplete policy evidence.

ARCHITECTURE

People Desk (browser)
    |
    |  GET /   POST /chat   POST /demos/{id}   POST /ask
    v
FastAPI  harborline.api
    |
    +-- /ask -------------------------------------------+
    |                                                    |
    +-- /chat, /demos                                    |
            |                                            |
            v                                            v
    Agent orchestrator                          retrieve path
    harborline.agent                            rewrite, embed, cosine
    regex intent, fixed workflow                rerank, guardrails
            |                                            |
            |  tools/list, tools/call                    |
            v                                            |
    MCP client  harborline.mcp_client                   |
    mcp-inproc in the web process                        |
    stdio for CLI and Cursor                             |
            |                                            |
            v                                            |
    MCP server  harborline.mcp_server (FastMCP)          |
            |                                            |
            +-- search_policy_documents                  |
            |   get_policy_section                       |
            |   check_policy_compliance ---------------> |
            |         RAG index                          |
            |         OpenAI embeddings + cosine cache   |
            |         corpus/ policies                   |
            |                                            |
            +-- lookup_employee_profile                  |
            |   check_pto_balance                        |
            |   lookup_benefits_status                   |
            |         mock HarborHub JSON in data/       |
            |                                            |
            +-- create_mock_hr_ticket                    |
                draft_hr_email                           |
                      session-only MOCK                  |
                                                         |
    After tools, if HARBORLINE_ANSWER_MODE=llm:          |
            v                                            |
    OpenAI chat API                                      |
    gpt-4o-mini, temperature 0, seed 42                  |
    writes the answer from the draft, snippets, trace    |

DEMO TASKS

Both demos are fixed in harborline/demos.py. POST /demos/{id} and POST /chat
with the same question and employee id run the same workflow. The client
opens the MCP bus, lists tools, then makes the calls below.

Demo 1. Remote eligibility (remote-emp-1008).
Question: Am I eligible for fully remote work living in Tacoma?
Employee: EMP-1008 (Alex Kim). Intent: remote_eligibility.

1. lookup_employee_profile
   employee_id=EMP-1008
   Reads the HarborHub profile, including location category and miles from
   the hub.
2. search_policy_documents
   query=remote hybrid hub 50 miles office days POL-RMT-003
   kind=policy
   Retrieves hub-rule passages.
3. get_policy_section
   policy_id=POL-RMT-003
   section=Location categories
   Loads that heading for the citation.
4. check_policy_compliance
   scenario=the original question
   employee_id=EMP-1008
   policy_id=POL-RMT-003
   Returns a verdict.

The draft states that Alex Kim is hub-coded, about 32 miles from Seattle, and
that fully remote work needs a People Operations reclass. The trace escalates
for that reclass. Citations come from corpus/03-remote-hybrid-work.md.

Demo 2. Benefits election (benefits-emp-1008).
Question: What medical plan and 401k deferral do I have?
Employee: EMP-1008. Intent: benefits.

1. search_policy_documents
   query=benefits 401k medical enrollment POL-BEN-006
   kind=policy
2. get_policy_section
   policy_id=POL-BEN-006
   section=US medical
3. get_policy_section
   policy_id=POL-BEN-006
   section=Retirement
4. lookup_benefits_status
   employee_id=EMP-1008
   Reads the mock election: HDHP medical plan and a 3 percent 401(k) deferral.

The draft quotes the 100 percent match on the first 4 percent deferred,
immediate vesting, and the HarborHub election row. Citations come from
corpus/06-employee-benefits.md.

With HARBORLINE_ANSWER_MODE=llm, both demos then send that draft, the
citations, and the tool summaries to the OpenAI chat API. The tool order
above stays the same.
""".strip()


def paginate(text: str) -> list[list[tuple[str, int, int]]]:
    lines = wrap(text, 88)
    pages: list[list[tuple[str, int, int]]] = []
    usable = 46
    chunks: list[list[str]] = []
    buf: list[str] = []
    for line in lines:
        buf.append(line)
        if len(buf) >= usable:
            chunks.append(buf)
            buf = []
    if buf:
        chunks.append(buf)
    total = len(chunks)
    for i, chunk in enumerate(chunks, start=1):
        items: list[tuple[str, int, int]] = []
        y = 750
        items.append(("Harborline  |  Design Documentation", 48, y))
        y -= 16
        items.append((f"Page {i} of {total}", 48, y))
        y -= 22
        for line in chunk:
            items.append((line, 48, y))
            y -= 14
        pages.append(items)
    return pages


DIAGRAM = """
People Desk (browser)
    |
    |  GET /   POST /chat   POST /demos/{id}   POST /ask
    v
FastAPI  harborline.api
    |
    +-- /ask -------------------------------------------+
    |                                                    |
    +-- /chat, /demos                                    |
            |                                            |
            v                                            v
    Agent orchestrator                          retrieve path
    harborline.agent                            rewrite, embed, cosine
    regex intent, fixed workflow                rerank, guardrails
            |                                            |
            |  tools/list, tools/call                    |
            v                                            |
    MCP client  harborline.mcp_client                   |
    mcp-inproc in the web process                        |
    stdio for CLI and Cursor                             |
            |                                            |
            v                                            |
    MCP server  harborline.mcp_server (FastMCP)          |
            |                                            |
            +-- search_policy_documents                  |
            |   get_policy_section                       |
            |   check_policy_compliance ---------------> |
            |         RAG index                          |
            |         OpenAI embeddings + cosine cache   |
            |         corpus/ policies                   |
            |                                            |
            +-- lookup_employee_profile                  |
            |   check_pto_balance                        |
            |   lookup_benefits_status                   |
            |         mock HarborHub JSON in data/       |
            |                                            |
            +-- create_mock_hr_ticket                    |
                draft_hr_email                           |
                      session-only MOCK                  |
                                                         |
    After tools, if HARBORLINE_ANSWER_MODE=llm:          |
            v                                            |
    OpenAI chat API                                      |
    gpt-4o-mini, temperature 0, seed 42                  |
    writes the answer from the draft, snippets, trace
""".strip("\n")


def _heading(doc, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def _para(doc, text: str, bold_lead: str | None = None) -> None:
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    if bold_lead:
        run = p.add_run(bold_lead)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def write_docx(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    doc.core_properties.title = "Design Documentation"
    doc.core_properties.subject = "Harborline design documentation"

    title = doc.add_heading("Design Documentation", 0)
    title.runs[0].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    sub = doc.add_paragraph()
    run = sub.add_run("Harborline")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    _para(
        doc,
        "This note records how Harborline answers employee questions: a hand-written "
        "orchestrator, one MCP server, and OpenAI embeddings with an in-process cosine "
        "index. The two People Desk demos follow fixed tool sequences.",
    )

    _heading(doc, "Design choices", 1)
    _para(
        doc,
        " harborline/agent.py classifies intent with regular expressions and then runs a "
        "fixed workflow. The model does not choose tools. That keeps HR paths repeatable: "
        "the same question always calls the same tools, and the trace lists each name, "
        "argument, and result. When HARBORLINE_ANSWER_MODE=llm and OPENAI_API_KEY is set, "
        "the chat model only rewrites the draft after the tools return. Refusals, missing "
        "records, and thin evidence stay on that draft.",
        bold_lead="Orchestration.",
    )
    _para(
        doc,
        " One FastMCP server, named harborline, exposes eight tools. Policy tools read the "
        "RAG index. Profile, PTO, and benefits tools read mock HarborHub JSON under data/. "
        "Ticket and email tools are session-only mocks and do not write data/tickets.json. "
        "The agent discovers tools with tools/list and calls only names from that list "
        "through tools/call. FastMCP builds each JSON Schema from the Python signature, so "
        "the schema stays tied to the function.",
        bold_lead="MCP server.",
    )
    _para(doc, "Tool schemas.", bold_lead="")
    schemas = [
        "search_policy_documents: query (required), employee_id, kind",
        "get_policy_section: policy_id (required), section",
        "lookup_employee_profile: employee_id",
        "check_pto_balance: employee_id",
        "lookup_benefits_status: employee_id",
        "create_mock_hr_ticket: topic, summary, employee_id, confirm",
        "draft_hr_email: employee_id, subject, body, confirm",
        "check_policy_compliance: scenario, employee_id, policy_id",
    ]
    for item in schemas:
        doc.add_paragraph(item, style="List Bullet")
    _para(
        doc,
        " The CLI and Cursor use stdio: the client spawns python -m harborline.mcp_server "
        "and speaks JSON-RPC on stdin and stdout, with no open port. The People Desk page, "
        "POST /chat, and POST /demos/{id} use mcp-inproc, the same FastMCP object inside "
        "the web process, so a Render container does not spawn a second Python process per "
        "request. Streamable HTTP on 127.0.0.1:8765 is available when an external client "
        "needs an HTTP MCP endpoint.",
        bold_lead="Transport.",
    )
    _para(
        doc,
        " Embeddings are OpenAI text-embedding-3-small (1536 dimensions) from the embeddings "
        "API. Ingest splits each handbook by heading. A section longer than 900 characters "
        "is cut into windows of 900 with 120 characters of overlap, so a policy fact stays "
        "with its heading and neighboring sentences still share context.",
        bold_lead="Embeddings and chunking.",
    )
    _para(
        doc,
        " Search embeds the rewritten query, takes a candidate pool of HARBORLINE_FETCH_K=20, "
        "then reranks to HARBORLINE_TOP_K=5. The blend is 0.65 cosine and 0.35 lexical "
        "overlap, and the reranker prefers distinct source files. The score floor is "
        "HARBORLINE_MIN_SCORE=0.22. Vectors are L2-normalized and cached under .cache/. "
        "Query time ranks that matrix in the process. tfidf remains the offline path used "
        "by pytest and CI.",
        bold_lead="Retrieval k and index.",
    )
    _para(
        doc,
        " A Docker image runs FastAPI with uvicorn on port 8000. The image copies corpus/, "
        "data/, and eval/ and does not download local embedding weights. On Render, set "
        "OPENAI_API_KEY and HARBORLINE_ANSWER_MODE=llm. The Dockerfile pins retrieve, so the "
        "service variable overrides it. The first search after a restart embeds the corpus "
        "and holds the matrix in that process. The cache does not survive a new deploy.",
        bold_lead="Deployment.",
    )
    _para(
        doc,
        " Out-of-corpus questions (stock tips, other employers, medical advice, and similar "
        "patterns) are refused. Empty retrieval or a top score under 0.22 returns a "
        "weak-match refusal. Missing employee ids ask for clarification. Missing HarborHub "
        "rows escalate. Mock ticket and email writes stay pending until confirm=true, and "
        "even then they exist only in that process. The model rewrite leaves refusals in "
        "place for an unavailable tool bus, a missing employee, or incomplete policy evidence.",
        bold_lead="Safety guardrails.",
    )

    _heading(doc, "Architecture", 1)
    for line in DIAGRAM.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    _heading(doc, "Demo tasks", 1)
    _para(
        doc,
        "Both demos are fixed in harborline/demos.py. POST /demos/{id} and POST /chat with "
        "the same question and employee id run the same workflow. The client opens the MCP "
        "bus, lists tools, then makes the calls below.",
    )
    _heading(doc, "Demo 1. Remote eligibility (remote-emp-1008)", 2)
    _para(doc, "Question: Am I eligible for fully remote work living in Tacoma?")
    _para(doc, "Employee: EMP-1008 (Alex Kim). Intent: remote_eligibility.")
    steps = [
        "1. lookup_employee_profile — employee_id=EMP-1008. Reads the HarborHub profile, including location category and miles from the hub.",
        "2. search_policy_documents — query “remote hybrid hub 50 miles office days POL-RMT-003”, kind=policy. Retrieves hub-rule passages.",
        "3. get_policy_section — policy_id=POL-RMT-003, section=Location categories. Loads that heading for the citation.",
        "4. check_policy_compliance — scenario is the original question, employee_id=EMP-1008, policy_id=POL-RMT-003. Returns a verdict.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Bullet")
    _para(
        doc,
        "The draft states that Alex Kim is hub-coded, about 32 miles from Seattle, and that "
        "fully remote work needs a People Operations reclass. The trace escalates for that "
        "reclass. Citations come from corpus/03-remote-hybrid-work.md.",
    )

    _heading(doc, "Demo 2. Benefits election (benefits-emp-1008)", 2)
    _para(doc, "Question: What medical plan and 401k deferral do I have?")
    _para(doc, "Employee: EMP-1008. Intent: benefits.")
    steps = [
        "1. search_policy_documents — query “benefits 401k medical enrollment POL-BEN-006”, kind=policy.",
        "2. get_policy_section — policy_id=POL-BEN-006, section=US medical.",
        "3. get_policy_section — policy_id=POL-BEN-006, section=Retirement.",
        "4. lookup_benefits_status — employee_id=EMP-1008. Reads the mock election: HDHP medical plan and a 3 percent 401(k) deferral.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Bullet")
    _para(
        doc,
        "The draft quotes the 100 percent match on the first 4 percent deferred, immediate "
        "vesting, and the HarborHub election row. Citations come from "
        "corpus/06-employee-benefits.md.",
    )
    _para(
        doc,
        "With HARBORLINE_ANSWER_MODE=llm, both demos then send that draft, the citations, "
        "and the tool summaries to the OpenAI chat API. The tool order above stays the same.",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    pdf = build_pdf(paginate(BODY))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(pdf)
    write_docx(DOCX)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    print(f"wrote {DOCX} ({DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
